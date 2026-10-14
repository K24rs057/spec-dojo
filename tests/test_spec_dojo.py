"""spec-dojo のテスト。AIは FakeLLM に置き換え、claude -p は呼ばない。

テスト仕様書の各行と対応するように、テスト名を日本語の説明つきで書く。
"""
import json
import sys
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import documents  # noqa: E402
import rules  # noqa: E402
from app import create_app  # noqa: E402
from jobs import JobRunner, Splitter  # noqa: E402
from llm import FakeLLM, LLMError  # noqa: E402
from prompts import MARKER  # noqa: E402
from store import Store  # noqa: E402

QUESTION = {
    "question": "部員が出欠を入れる場面は?",
    "options": [{"id": "A", "label": "前日にスマホで", "advice": "スマホ前提になります"},
                {"id": "B", "label": "週の初めにPCで", "advice": "入れ忘れに注意"},
                {"id": "C", "label": "幹事が代わりに", "advice": "目的と食い違うかも"}],
    "hint": "自分が部員だったらを想像する",
    "recommended": ["A"],
    "terms": [{"term": "要件定義", "explanation": "作る前に何を作るか決めること"}],
}


def reply(text, obj):
    return f"{text}\n{MARKER}\n{json.dumps(obj, ensure_ascii=False)}"


@pytest.fixture
def env(tmp_path):
    store = Store(":memory:")
    llm = FakeLLM([])
    runner = JobRunner(llm)
    client = TestClient(create_app(store, runner, tmp_path / "output"))
    pid = client.post("/api/projects", json={"name": "サークルの出欠アプリ"}).json()["id"]
    return {"store": store, "llm": llm, "runner": runner, "client": client, "pid": pid, "tmp": tmp_path}


def wait_job(runner, job_id):
    job = runner.get(job_id)
    for _ in range(200):
        if job.done:
            return job
        time.sleep(0.01)
    raise AssertionError("ジョブが終わらない")


def result_of(job):
    kind, data = job.events[-1][1], job.events[-1][2]
    return kind, data


# ---- ①曖昧語辞書・秘密情報(AIなし) ----
def test_曖昧語を見つける():
    ws = rules.lint("なるべく見やすく表示する")
    assert {w["word"] for w in ws} >= {"なるべく", "見やすく"}


def test_曖昧語がなければ警告しない():
    assert rules.lint("1秒以内に一覧を表示する") == []


@pytest.mark.parametrize("text", ["sk-abcdefghijk123", "password = hunter2", "AIzaSyA1234567890abc"])
def test_秘密情報らしき文字列を見つける(text):
    assert rules.contains_secret(text)


def test_秘密情報を含む依頼はAIに送らない(env):
    r = env["client"].post("/api/jobs", json={"project_id": env["pid"], "kind": "chat",
                                               "message": "キーは sk-abcdefghijk123 です"})
    assert r.status_code == 400
    assert env["llm"].prompts == []


# ---- 決定と履歴 ----
def test_決定すると履歴が残る(env):
    s, pid = env["store"], env["pid"]
    s.decide(pid, "requirements.purpose", "出欠を一目で分かるようにする")
    s.decide(pid, "requirements.purpose", "出欠を30秒で確認できるようにする", reason="数字にした")
    h = [x for x in s.history(pid) if x["item_key"] == "requirements.purpose"]
    assert [x["new_value"] for x in h] == ["出欠を一目で分かるようにする", "出欠を30秒で確認できるようにする"]
    assert h[1]["old_value"] == "出欠を一目で分かるようにする"


def test_前提を変えると依存している決定に要確認が付く(env):
    s, pid = env["store"], env["pid"]
    s.decide(pid, "requirements.purpose", "目的1")
    s.decide(pid, "requirements.users", "幹事と部員")
    marked = s.decide(pid, "requirements.purpose", "目的2")
    assert marked == ["requirements.users"]
    assert s.decision(pid, "requirements.users")["status"] == "needs_check"


def test_未決定の項目には要確認を付けない(env):
    s, pid = env["store"], env["pid"]
    s.decide(pid, "requirements.purpose", "目的1")
    assert s.decide(pid, "requirements.purpose", "目的2") == []


def test_このままでいいで要確認が外れる(env):
    s, pid, c = env["store"], env["pid"], env["client"]
    s.decide(pid, "requirements.purpose", "目的1")
    s.decide(pid, "requirements.users", "部員")
    s.decide(pid, "requirements.purpose", "目的2")
    c.post(f"/api/projects/{pid}/decisions/requirements.users/keep")
    assert s.decision(pid, "requirements.users")["status"] == "decided"


def test_履歴の書き込みに失敗したら決定もなかったことになる(env, monkeypatch):
    s, pid = env["store"], env["pid"]
    s.conn.execute("CREATE TRIGGER boom BEFORE INSERT ON decision_history BEGIN SELECT RAISE(ABORT, 'x'); END")
    with pytest.raises(Exception):
        s.decide(pid, "requirements.purpose", "目的")
    assert s.decision(pid, "requirements.purpose")["status"] == "undecided"


# ---- 工程の完了 ----
def test_要確認が残っていると工程を完了できない(env):
    s, pid, c = env["store"], env["pid"], env["client"]
    for it in rules.items("requirements"):
        s.decide(pid, it["key"], "内容")
    s.decide(pid, "requirements.purpose", "変えた")
    r = c.post(f"/api/projects/{pid}/phases/requirements/complete")
    assert r.status_code == 409 and "要確認" in r.json()["detail"]


def test_全部決まれば次の工程に進む(env):
    s, pid, c = env["store"], env["pid"], env["client"]
    for it in rules.items("requirements"):
        s.decide(pid, it["key"], "内容")
    r = c.post(f"/api/projects/{pid}/phases/requirements/complete")
    assert r.json()["next_phase"] == "concept"


# ---- アーカイブ ----
def test_しまったプロジェクトは一覧から消えて元に戻せる(env):
    c, pid = env["client"], env["pid"]
    c.post(f"/api/projects/{pid}/archive")
    assert c.get("/api/projects").json() == []
    assert len(c.get("/api/projects?archived=1").json()) == 1
    c.post(f"/api/projects/{pid}/restore")
    assert len(c.get("/api/projects").json()) == 1


# ---- AIのジョブ ----
def test_区切り行が2つの断片にまたがってもJSONを取りこぼさない():
    sp = Splitter()
    shown = sp.feed("説明です。\n<<<SPEC_DO") + sp.feed('JO_JSON>>>\n{"a": 1}')
    assert "<<<" not in shown and shown.strip() == "説明です。"
    assert json.loads(sp.json_part) == {"a": 1}


def test_質問を作ると説明文だけが流れて質問が保存される(env):
    env["llm"].replies.append(reply("次は利用場面です。", QUESTION))
    r = env["client"].post("/api/jobs", json={"project_id": env["pid"], "kind": "question",
                                               "item_key": "requirements.scene"})
    job = wait_job(env["runner"], r.json()["job_id"])
    texts = "".join(d for _, k, d in job.events if k == "text")
    assert MARKER not in texts and "次は利用場面です。" in texts
    kind, data = result_of(job)
    assert kind == "result" and data["turn"]["question"]["question"] == QUESTION["question"]
    assert env["store"].glossary(env["pid"])[0]["term"] == "要件定義"


def test_答えていない質問があれば作り直さない(env):
    env["llm"].replies.append(reply("説明", QUESTION))
    c = env["client"]
    r = c.post("/api/jobs", json={"project_id": env["pid"], "kind": "question", "item_key": "requirements.scene"})
    wait_job(env["runner"], r.json()["job_id"])
    r2 = c.post("/api/jobs", json={"project_id": env["pid"], "kind": "question", "item_key": "requirements.scene"})
    assert "turn" in r2.json() and len(env["llm"].prompts) == 1


def test_JSONの形が崩れていたら1回だけやり直す(env):
    env["llm"].replies += [reply("説明", {"question": "形が違う"}), reply("説明", QUESTION)]
    r = env["client"].post("/api/jobs", json={"project_id": env["pid"], "kind": "question",
                                               "item_key": "requirements.scene"})
    job = wait_job(env["runner"], r.json()["job_id"])
    assert any(k == "reset" for _, k, _ in job.events)
    assert result_of(job)[0] == "result"


def test_2回続けて崩れたらエラーを返す(env):
    env["llm"].replies += ["区切りなしの返事", "区切りなしの返事"]
    r = env["client"].post("/api/jobs", json={"project_id": env["pid"], "kind": "question",
                                               "item_key": "requirements.scene"})
    job = wait_job(env["runner"], r.json()["job_id"])
    assert result_of(job)[0] == "error"


def test_AIが失敗してもそれまでの決定は消えない(env):
    s, pid = env["store"], env["pid"]
    s.decide(pid, "requirements.purpose", "出欠を30秒で確認できるようにする")
    env["llm"].replies.append(LLMError("AIの返事が180秒以内に終わらなかったので打ち切りました"))
    r = env["client"].post("/api/jobs", json={"project_id": pid, "kind": "question", "item_key": "requirements.users"})
    job = wait_job(env["runner"], r.json()["job_id"])
    kind, data = result_of(job)
    assert kind == "error" and "180秒" in data["message"]
    assert s.decision(pid, "requirements.purpose")["value"] == "出欠を30秒で確認できるようにする"


def test_選択肢を1つ選んだら先回りのアドバイスを待ち時間なしで返す(env):
    env["llm"].replies.append(reply("説明", QUESTION))
    c = env["client"]
    r = c.post("/api/jobs", json={"project_id": env["pid"], "kind": "question", "item_key": "requirements.scene"})
    turn = result_of(wait_job(env["runner"], r.json()["job_id"]))[1]["turn"]
    a = c.post("/api/jobs", json={"project_id": env["pid"], "kind": "advice", "turn_id": turn["id"], "selected": ["B"]})
    assert a.json()["advice"]["text"] == "入れ忘れに注意"
    assert len(env["llm"].prompts) == 1  # AIは追加で呼んでいない


def test_2つ以上選んだらAIを追加で呼ぶ(env):
    env["llm"].replies += [reply("説明", QUESTION),
                           reply("両方は成り立ちます", {"value": "前日スマホ、まとめてPCも可", "contradiction": ""})]
    c = env["client"]
    r = c.post("/api/jobs", json={"project_id": env["pid"], "kind": "question", "item_key": "requirements.scene"})
    turn = result_of(wait_job(env["runner"], r.json()["job_id"]))[1]["turn"]
    a = c.post("/api/jobs", json={"project_id": env["pid"], "kind": "advice", "turn_id": turn["id"], "selected": ["A", "B"]})
    data = result_of(wait_job(env["runner"], a.json()["job_id"]))[1]
    assert data["advice"]["value"] == "前日スマホ、まとめてPCも可"
    assert len(env["llm"].prompts) == 2


def test_自由チャットで知らない項目キーを返されたらやり直す(env):
    env["llm"].replies += [reply("返事", {"candidates": [{"item_key": "nope", "value": "x"}]}),
                           reply("返事", {"candidates": [{"item_key": "requirements.scene", "value": "前日スマホ"}]})]
    r = env["client"].post("/api/jobs", json={"project_id": env["pid"], "kind": "chat", "message": "スマホで前日に"})
    kind, data = result_of(wait_job(env["runner"], r.json()["job_id"]))
    assert kind == "result" and data["candidates"][0]["item_key"] == "requirements.scene"


# ---- レビュー ----
def test_レビューは曖昧語辞書とAIの指摘を保存してレビュー済みにする(env):
    s, pid = env["store"], env["pid"]
    s.decide(pid, "requirements.purpose", "なるべく早く出欠を確認する")
    env["llm"].replies.append(reply("全体に良いです", {"findings": [
        {"item_key": "requirements.purpose", "viewpoint_id": "V-06", "message": "達成を確かめる方法がない"}]}))
    r = env["client"].post("/api/jobs", json={"project_id": pid, "kind": "review"})
    kind, data = result_of(wait_job(env["runner"], r.json()["job_id"]))
    layers = sorted(f["layer"] for f in data["findings"])
    assert kind == "result" and layers[0] == 1 and 2 in layers
    assert s.unreviewed(pid) == []


# ---- 資料と版番号 ----
def test_版は出力のたびに上がり内容が同じなら上がらない(env):
    s, pid = env["store"], env["pid"]
    p = s.get_project(pid)
    s.decide(pid, "requirements.purpose", "目的1")
    _, r1 = documents.export(s, p, env["tmp"] / "out", ("md",))
    _, r2 = documents.export(s, p, env["tmp"] / "out", ("md",))
    s.decide(pid, "requirements.purpose", "目的2")
    _, r3 = documents.export(s, p, env["tmp"] / "out", ("md",))
    v = lambda rs: next(x["version"] for x in rs if x["doc_type"] == "requirements")
    assert (v(r1), v(r2), v(r3)) == ("0.1", "0.1", "0.2")


def test_工程を完了すると1_00になりその後は1_01(env):
    s, pid = env["store"], env["pid"]
    for it in rules.items("requirements"):
        s.decide(pid, it["key"], "内容")
    s.set_phase(pid, "concept")
    p = s.get_project(pid)
    _, r1 = documents.export(s, p, env["tmp"] / "out", ("md",))
    s.decide(pid, "requirements.purpose", "変えた")
    _, r2 = documents.export(s, p, env["tmp"] / "out", ("md",))
    v = lambda rs: next(x["version"] for x in rs if x["doc_type"] == "requirements")
    assert (v(r1), v(r2)) == ("1.00", "1.01")


def test_Excelに表紙と改版履歴と一覧と疑問点が入る(env):
    s, pid, c = env["store"], env["pid"], env["client"]
    s.decide(pid, "requirements.purpose", "目的")
    r = c.post(f"/api/projects/{pid}/exports", json={"formats": ["xlsx", "md"]}).json()
    xlsx = next(Path(r["folder"]).glob("requirements-v0.1.xlsx"))
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["表紙", "改版履歴", "要件定義一覧", "疑問点"]
    assert "下書き" in wb["表紙"]["C5"].value
    assert (Path(r["folder"]) / "implementation-guide-v0.1.md").exists()


def test_zipをダウンロードできる(env):
    c, pid = env["client"], env["pid"]
    r = c.post(f"/api/projects/{pid}/exports", json={"formats": ["md"], "zip": True}).json()
    z = c.get(r["zip_url"])
    assert z.status_code == 200 and z.content[:2] == b"PK"
