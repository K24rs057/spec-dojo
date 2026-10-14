"""決定事項から資料(Markdown / Excel)を作り、版番号を決める(F9〜F12)。

版番号のルール(基本設計書 1 疑問点9):
- 工程が完了する前は、出力のたびに 0.1, 0.2 … と上がる(下書き)
- 工程を完了したら 1.00。そのあと内容を変えて出力したら 1.01, 1.02 …
- 前回の出力から内容が変わっていなければ、版は上げない
"""
import hashlib
import re
import shutil
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import rules

DOC_TYPES = {"requirements": "要件定義書", "concept": "構想設計書", "basic": "基本設計書",
             "guide": "実装指示書"}
FILE_NAMES = {"requirements": "requirements", "concept": "concept-design", "basic": "basic-design",
              "guide": "implementation-guide"}
STATUS_LABEL = {"undecided": "未決定", "decided": "決定", "needs_check": "要確認"}


def phase_completed(project, doc_type):
    order = rules.phase_keys() + ["done"]
    target = "basic" if doc_type == "guide" else doc_type
    return order.index(project["current_phase"]) > order.index(target)


def body_markdown(store, project, doc_type):
    """版番号を含まない本文(内容の指紋はこれから計算する)。"""
    decs = store.decisions(project["id"])
    if doc_type == "guide":
        return guide_markdown(project, decs)
    phase = next(p for p in rules.phases() if p["key"] == doc_type)
    lines = [f"| 項目 | 内容 | 理由 | 出どころ | 状態 |", "|---|---|---|---|---|"]
    for it in phase["items"]:
        d = decs[it["key"]]
        cell = lambda s: (s or "").replace("|", "/").replace("\n", " ")
        lines.append(f"| {it['label']} | {cell(d['value']) or '—'} | {cell(d['reason'])} | "
                     f"{cell(d['source'])} | {STATUS_LABEL[d['status']]} |")
    open_items = [it["label"] for it in phase["items"] if decs[it["key"]]["status"] != "decided"]
    lines += ["", "## 疑問点(まだ決まっていないこと)", ""]
    lines += [f"- {x}" for x in open_items] or ["- なし"]
    findings = [f for f in store.findings(project["id"]) if rules.phase_of(f["item_key"]) == doc_type]
    if findings:
        lines += ["", "## レビューの指摘(未対応)", ""]
        lines += [f"- [{f['viewpoint_id'] or '辞書'}] {rules.item(f['item_key'])['label']}:{f['message']}"
                  for f in findings]
    return "\n".join(lines) + "\n"


def guide_markdown(project, decs):
    def val(key):
        d = decs[key]
        return d["value"] if d["status"] != "undecided" else "(未決定)"

    out = [
        "## この指示書の読み方",
        "",
        "これは、AIコーディングツール(Claude Code、Codex、Cursor など)に渡すための実装指示書です。",
        "仕様道場で、要件定義・構想設計・基本設計を本人が決めた内容をまとめています。",
        "ここに書かれていることは本人が決めたことなので、勝手に変えず、迷ったら本人に確認してください。",
        "",
        "## AIに任せる範囲(詳細設計)",
        "",
        "- ファイルとフォルダの分け方、関数やクラスの設計",
        "- ライブラリの細かい選び方(下の「制約」に反しない範囲で)",
        "- 変数名、エラーメッセージの細かい文言",
        "- テストの書き方",
        "",
        "上の範囲で決めたことは、作ったあとに本人に一覧で報告してください。",
        "",
    ]
    for p in rules.phases():
        out += [f"## {p['label']}", ""]
        for it in p["items"]:
            out.append(f"### {it['label']}")
            out.append("")
            out.append(val(it["key"]))
            out.append("")
    out += ["## 必ず守ること", "",
            f"- 対象外:{val('requirements.out_of_scope')}",
            f"- 制約:{val('requirements.constraints')}",
            f"- 非機能要件:{val('requirements.nonfunctional')}", ""]
    return "\n".join(out)


def content_hash(text):
    return hashlib.sha256(text.encode("utf-8")).hexdigest()[:16]


def next_version(store, project, doc_type, h):
    last = store.last_export(project["id"], doc_type)
    if last and last["content_hash"] == h:
        return last["version"], False
    history = store.exports(project["id"], doc_type)
    if phase_completed(project, doc_type):
        released = [e for e in history if not e["version"].startswith("0.")]
        if not released:
            return "1.00", True
        major, minor = released[-1]["version"].split(".")
        return f"{major}.{int(minor) + 1:02d}", True
    drafts = [e for e in history if e["version"].startswith("0.")]
    return f"0.{len(drafts) + 1}", True


def render_markdown(store, project, doc_type, version):
    title = DOC_TYPES[doc_type]
    draft = version.startswith("0.")
    head = [f"# {title}:{project['name']}", "",
            f"ver.{version}{'(下書き)' if draft else ''} / {date.today().isoformat()} / 仕様道場で作成", ""]
    hist = store.exports(project["id"], doc_type)
    rows = ["## 改版履歴", "", "| 版 | 日付 |", "|---|---|"]
    rows += [f"| {e['version']} | {e['exported_at'][:10]} |" for e in hist]
    return "\n".join(head) + body_markdown(store, project, doc_type) + "\n" + "\n".join(rows) + "\n"


def preview(store, project, doc_type):
    body = body_markdown(store, project, doc_type)
    version, changed = next_version(store, project, doc_type, content_hash(body))
    return {"doc_type": doc_type, "title": DOC_TYPES[doc_type], "version": version,
            "draft": version.startswith("0."), "markdown": render_markdown(store, project, doc_type, version)}


def safe_name(name):
    return re.sub(r'[\\/:*?"<>|\s]+', "_", name).strip("_") or "project"


def export(store, project, out_root, formats=("md", "xlsx")):
    """全資料を output/<id>_<名前>/ に書き出す。戻り値は書き出したフォルダと各資料の版。"""
    folder = Path(out_root) / f"{project['id']}_{safe_name(project['name'])}"
    folder.mkdir(parents=True, exist_ok=True)
    results = []
    for doc_type in DOC_TYPES:
        body = body_markdown(store, project, doc_type)
        h = content_hash(body)
        version, changed = next_version(store, project, doc_type, h)
        if changed:
            store.add_export(project["id"], doc_type, version, h)
        base = f"{FILE_NAMES[doc_type]}-v{version}"
        md = render_markdown(store, project, doc_type, version)
        if "md" in formats:
            (folder / f"{base}.md").write_text(md, encoding="utf-8")
        if "xlsx" in formats and doc_type != "guide":
            write_xlsx(store, project, doc_type, version, folder / f"{base}.xlsx")
        results.append({"doc_type": doc_type, "version": version, "new_version": changed})
    return folder, results


def make_zip(folder):
    return Path(shutil.make_archive(str(folder), "zip", root_dir=folder))


# ---- Excel(授業の要件定義書と同じ並び:表紙 → 改版履歴 → 一覧 → 疑問点) ----
_thin = Side(style="thin", color="000000")
_BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_HEAD = PatternFill("solid", fgColor="DDEBF7")
_WRAP = Alignment(wrap_text=True, vertical="top")
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_FONT = "Yu Gothic"


def _f(size=10, bold=False):
    return Font(name=_FONT, size=size, bold=bold)


def _table(ws, top, headers, rows, widths):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=top, column=i, value=h)
        c.font, c.fill, c.border, c.alignment = _f(10, True), _HEAD, _BOX, _CENTER
    for r, row in enumerate(rows, top + 1):
        for i, v in enumerate(row, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font, c.border, c.alignment = _f(10), _BOX, _WRAP
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w


def write_xlsx(store, project, doc_type, version, path):
    title = DOC_TYPES[doc_type]
    decs = store.decisions(project["id"])
    phase = next(p for p in rules.phases() if p["key"] == doc_type)
    draft = version.startswith("0.")
    done = sum(1 for it in phase["items"] if decs[it["key"]]["status"] == "decided")

    wb = Workbook()
    ws = wb.active
    ws.title = "表紙"
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGH", (4, 4, 14, 14, 14, 14, 14, 14)):
        ws.column_dimensions[col].width = w
    ws.merge_cells("C5:H5")
    ws["C5"] = title + ("(下書き)" if draft else "")
    ws["C5"].font, ws["C5"].alignment = _f(26, True), _CENTER
    ws.row_dimensions[5].height = 46
    for r, (k, v) in enumerate((("システム名称", project["name"]), ("版", f"ver.{version}"),
                                ("日付", date.today().strftime("%Y年%m月%d日")),
                                ("進み具合", f"{done}/{len(phase['items'])}")), 8):
        ws.cell(row=r, column=3, value=k).font = _f(12, True)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)
        ws.cell(row=r, column=4, value=v).font = _f(12)
    ws.merge_cells("C14:E14")
    ws.merge_cells("F14:H14")
    ws["C14"], ws["F14"] = "作成", "承認"
    for col, label in zip("CDEFGH", ("氏名", "日付", "印") * 2):
        ws[f"{col}15"] = label
    for row in (14, 15, 16):
        for col in "CDEFGH":
            c = ws[f"{col}{row}"]
            c.border, c.alignment, c.font = _BOX, _CENTER, _f(10, row != 16)
    ws.row_dimensions[16].height = 42

    ws = wb.create_sheet("改版履歴")
    hist = store.exports(project["id"], doc_type)
    rows = [[i, e["version"], "下書き" if e["version"].startswith("0.") else "確定", e["exported_at"][:10]]
            for i, e in enumerate(hist, 1)]
    _table(ws, 1, ["項番", "改定番号", "区分", "日付"], rows, [6, 10, 10, 14])

    ws = wb.create_sheet(f"{phase['label']}一覧")
    rows = [[i, it["label"], decs[it["key"]]["value"] or "未決定", decs[it["key"]]["reason"] or "",
             decs[it["key"]]["source"] or "", STATUS_LABEL[decs[it["key"]]["status"]]]
            for i, it in enumerate(phase["items"], 1)]
    _table(ws, 1, ["No", "項目", "内容", "理由", "出どころ", "状態"], rows, [5, 16, 60, 30, 16, 8])

    ws = wb.create_sheet("疑問点")
    rows = [[it["label"], "まだ決まっていない" if decs[it["key"]]["status"] == "undecided" else "要確認(前提が変わった)",
             it["guide"]] for it in phase["items"] if decs[it["key"]]["status"] != "decided"]
    rows += [[rules.item(f["item_key"])["label"], f"レビュー指摘 {f['viewpoint_id'] or '曖昧語'}", f["message"]]
             for f in store.findings(project["id"]) if rules.phase_of(f["item_key"]) == doc_type]
    _table(ws, 1, ["項目", "種類", "内容・次に決めること"], rows or [["なし", "", ""]], [16, 24, 80])

    for s in wb.worksheets:
        s.page_setup.orientation = "landscape"
    wb.save(path)
